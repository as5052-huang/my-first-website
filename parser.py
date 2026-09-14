from __future__ import annotations

import re
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from pypdf import PdfReader

MAX_CHARS = 28000
MAX_BYTES = 20 * 1024 * 1024
_CJK_RE = re.compile(r"[\u4e00-\u9fff]")
_CID_RE = re.compile(r"\(cid:\d+\)", re.I)
_GARBLED_RE = re.compile(r"[\ufffd\u0000]")
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class ResumeParseError(ValueError):
    pass


def extract_resume_text(filename: str, data: bytes) -> str:
    if not data:
        raise ResumeParseError("文件为空，请重新上传。")
    if len(data) > MAX_BYTES:
        raise ResumeParseError("文件过大（超过 20MB），请压缩或另存后再上传。")

    kind = _sniff_kind(filename, data)
    if kind == "pdf":
        text = _from_pdf(data)
    elif kind == "docx":
        text = _from_docx(data)
    elif kind == "doc":
        text = _from_legacy_doc(data)
    elif kind == "rtf":
        text = _from_rtf(data)
    elif kind == "txt":
        text = _from_txt(data)
    else:
        raise ResumeParseError("仅支持 PDF、DOCX、DOC、RTF、TXT。请用 Word 另存为 .docx 后再试。")

    text = _normalize(text)
    if _looks_garbled(text):
        raise ResumeParseError("文字提取后乱码较多，请将简历另存为「.docx」或可复制文字的 PDF 后再上传。")
    if not text.strip():
        raise ResumeParseError(
            "未能提取到文字。若是扫描件/图片型 PDF，请先 OCR 或直接粘贴文本；"
            "若是旧版 .doc，请用 Word/WPS 另存为 .docx。"
        )
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS]
    return text


def _sniff_kind(filename: str, data: bytes) -> str:
    head = data[:8]
    suffix = Path(filename).suffix.lower()
    stripped = data.lstrip()

    if head.startswith(b"%PDF") or suffix == ".pdf":
        return "pdf"
    if head.startswith(b"PK") or suffix == ".docx":
        if _zip_has_word_xml(data):
            return "docx"
        if suffix == ".docx":
            return "docx"
        return "txt"
    if head.startswith(b"\xd0\xcf\x11\xe0") or suffix == ".doc":
        if _zip_has_word_xml(data):
            return "docx"
        return "doc"
    if stripped.startswith(b"{\\rtf") or suffix == ".rtf":
        return "rtf"
    if suffix == ".txt":
        return "txt"
    if stripped.startswith(b"<?xml") and b"wordprocessingml" in data[:4000].lower():
        return "docx"
    if suffix in {".pdf", ".docx", ".doc", ".rtf", ".txt"}:
        return suffix.lstrip(".")
    return ""


def _zip_has_word_xml(data: bytes) -> bool:
    try:
        with zipfile.ZipFile(BytesIO(data)) as zf:
            names = set(zf.namelist())
        return "word/document.xml" in names or "[Content_Types].xml" in names
    except zipfile.BadZipFile:
        return False


def _from_pdf(data: bytes) -> str:
    candidates: list[str] = []
    extractor_errors: dict[str, str] = {}

    for name, extractor in [
        ("pymupdf",   _pdf_pymupdf),
        ("pdfminer",  _pdf_pdfminer),
        ("pypdf",     _pdf_pypdf),
    ]:
        try:
            extracted = extractor(data)
        except Exception as exc:  # 记录每个提取器的具体错误，便于调试
            extractor_errors[name] = str(exc)
            continue
        if extracted and extracted.strip():
            candidates.append(extracted)

    if candidates:
        return _best_text(candidates)

    # 所有提取器均失败，给出具体原因
    if not extractor_errors:
        return ""
    hint_parts = [f"{n}: {e}" for n, e in extractor_errors.items()]
    hint = "；".join(hint_parts)
    if "password" in hint.lower() or "encrypt" in hint.lower():
        raise ResumeParseError("PDF 已加密，请先解除密码再上传。") from None
    if "corrupt" in hint.lower() or "invalid" in hint.lower() or "parse" in hint.lower():
        raise ResumeParseError(f"PDF 文件损坏或格式异常：{hint}") from None
    raise ResumeParseError(f"PDF 文字提取失败（尝试了 pymupdf、pdfminer、pypdf 均未提取到内容）：{hint}")


def _pdf_pypdf(data: bytes) -> str:
    reader = PdfReader(BytesIO(data), strict=False)
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ResumeParseError("PDF 已加密，请先解除密码再上传。") from exc
    pages: list[str] = []
    for page in reader.pages:
        chunk = ""
        try:
            chunk = page.extract_text(layout_mode="layout") or ""
        except TypeError:
            chunk = page.extract_text() or ""
        except Exception:
            try:
                chunk = page.extract_text() or ""
            except Exception:
                chunk = ""
        pages.append(chunk)
    return "\n".join(pages)


def _pdf_pdfminer(data: bytes) -> str:
    from pdfminer.high_level import extract_text

    return extract_text(BytesIO(data)) or ""


def _pdf_pymupdf(data: bytes) -> str:
    import fitz

    doc = fitz.open(stream=data, filetype="pdf")
    try:
        if doc.needs_pass:
            if not doc.authenticate(""):
                raise ResumeParseError("PDF 已加密，请先解除密码再上传。")
        pages: list[str] = []
        for page in doc:
            text = page.get_text("text") or ""
            if len(text.strip()) < 8:
                blocks = page.get_text("blocks") or []
                blocks = [b for b in blocks if isinstance(b, (list, tuple)) and len(b) >= 5]
                blocks.sort(key=lambda b: (round(float(b[1]), 1), round(float(b[0]), 1)))
                text = "\n".join(str(b[4]).strip() for b in blocks if str(b[4]).strip())
            pages.append(text)
        return "\n".join(pages)
    finally:
        doc.close()


def _from_docx(data: bytes) -> str:
    candidates: list[str] = []
    stripped = data.lstrip()
    if stripped.startswith(b"<?xml"):
        candidates.append(_xml_plain_text(data))
    else:
        try:
            candidates.append(_docx_via_python_docx(data))
        except Exception:
            pass
        try:
            candidates.append(_docx_via_xml_zip(data))
        except Exception:
            pass
    return _best_text([c for c in candidates if c and c.strip()])


def _docx_via_python_docx(data: bytes) -> str:
    doc = Document(BytesIO(data))
    parts: list[str] = []
    parts.extend(_iter_paragraphs(doc.paragraphs))
    parts.extend(_iter_tables(doc.tables))
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            try:
                parts.extend(_iter_paragraphs(header.paragraphs))
                parts.extend(_iter_tables(header.tables))
            except Exception:
                continue
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            try:
                parts.extend(_iter_paragraphs(footer.paragraphs))
                parts.extend(_iter_tables(footer.tables))
            except Exception:
                continue
    return "\n".join(p for p in parts if p)


def _iter_paragraphs(paragraphs) -> list[str]:
    return [p.text.strip() for p in paragraphs if getattr(p, "text", "").strip()]


def _iter_tables(tables) -> list[str]:
    lines: list[str] = []
    for table in tables:
        for row in table.rows:
            seen: set[int] = set()
            cells: list[str] = []
            for cell in row.cells:
                marker = id(cell._tc)
                if marker in seen:
                    continue
                seen.add(marker)
                text = " ".join(cell.text.split())
                if text:
                    cells.append(text)
                try:
                    lines.extend(_iter_tables(cell.tables))
                except Exception:
                    pass
            if cells:
                lines.append(" | ".join(cells))
    return lines


def _docx_via_xml_zip(data: bytes) -> str:
    with zipfile.ZipFile(BytesIO(data)) as zf:
        names = [
            n
            for n in zf.namelist()
            if n.startswith("word/")
            and n.endswith(".xml")
            and not n.endswith(".xml.rels")
            and "/_rels/" not in n
        ]
        preferred = [n for n in names if n == "word/document.xml"]
        others = sorted(n for n in names if n != "word/document.xml")
        chunks: list[str] = []
        for name in preferred + others:
            try:
                xml_bytes = zf.read(name)
            except KeyError:
                continue
            text = _xml_plain_text(xml_bytes)
            if text.strip():
                chunks.append(text)
    return "\n".join(chunks)


def _xml_plain_text(xml_bytes: bytes) -> str:
    try:
        from lxml import etree

        root = etree.fromstring(
            xml_bytes,
            parser=etree.XMLParser(recover=True, huge_tree=True, resolve_entities=False),
        )
        nodes = root.iter()
    except Exception:
        try:
            root = ET.fromstring(xml_bytes)
            nodes = root.iter()
        except ET.ParseError:
            return ""

    lines: list[str] = []
    current: list[str] = []
    for node in nodes:
        tag = _local_tag(getattr(node, "tag", ""))
        if tag in {"t", "delText"}:
            text = node.text or ""
            if text:
                current.append(text)
        elif tag in {"tab"}:
            current.append("\t")
        elif tag in {"br", "cr"}:
            current.append("\n")
        elif tag == "p":
            if current:
                line = "".join(current).strip()
                if line:
                    lines.append(line)
                current = []
    if current:
        line = "".join(current).strip()
        if line:
            lines.append(line)
    if not lines:
        fallback: list[str] = []
        for node in root.iter():
            if _local_tag(getattr(node, "tag", "")) in {"t", "delText"} and node.text:
                fallback.append(node.text)
        return _normalize("".join(fallback))
    return "\n".join(lines)


def _local_tag(tag: str) -> str:
    if not isinstance(tag, str):
        return ""
    if "}" in tag:
        return tag.rsplit("}", 1)[-1]
    return tag


def _from_legacy_doc(data: bytes) -> str:
    if _zip_has_word_xml(data):
        return _from_docx(data)
    salvaged = _salvage_binary_text(data)
    if _score_text(salvaged) >= 12:
        return salvaged
    raise ResumeParseError("检测到旧版 .doc（Word 97-2003）。请用 Word 或 WPS 另存为 .docx 后再上传。")


def _from_rtf(data: bytes) -> str:
    raw = _from_txt(data)
    raw = re.sub(r"\\'[0-9a-fA-F]{2}", "", raw)
    raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", raw)
    raw = raw.replace("{", " ").replace("}", " ")
    return _normalize(raw)


def _from_txt(data: bytes) -> str:
    if data.startswith(b"\xff\xfe"):
        return data.decode("utf-16-le", errors="replace")
    if data.startswith(b"\xfe\xff"):
        return data.decode("utf-16-be", errors="replace")
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig")

    try:
        from charset_normalizer import from_bytes

        best = from_bytes(data).best()
        if best is not None:
            decoded = str(best)
            if not _looks_garbled(decoded):
                return decoded
    except Exception:
        pass

    ranked: list[str] = []
    for enc in ("utf-8", "gb18030", "gbk", "gb2312", "big5", "utf-16", "latin-1"):
        try:
            ranked.append(data.decode(enc))
        except UnicodeDecodeError:
            continue
    if not ranked:
        return data.decode("utf-8", errors="replace")
    return _best_text(ranked)


def _salvage_binary_text(data: bytes) -> str:
    chunks: list[str] = []
    try:
        utf16 = data.decode("utf-16-le", errors="ignore")
        chunks.append(_keep_readable_runs(utf16))
    except Exception:
        pass
    try:
        chunks.append(_keep_readable_runs(data.decode("gb18030", errors="ignore")))
    except Exception:
        pass
    return _best_text(chunks)


def _keep_readable_runs(text: str) -> str:
    cleaned = _CONTROL_RE.sub(" ", text)
    runs = re.findall(r"[\u4e00-\u9fffA-Za-z0-9，。、；：？！“”‘’（）()\\-_.@/ ]{6,}", cleaned)
    return _normalize("\n".join(runs))


def _best_text(candidates: list[str]) -> str:
    usable = [_normalize(c) for c in candidates if c and str(c).strip()]
    if not usable:
        return ""
    return max(usable, key=_score_text)


def _score_text(text: str) -> float:
    if not text or not text.strip():
        return -100.0
    cjk = len(_CJK_RE.findall(text))
    cid = len(_CID_RE.findall(text))
    garbled = len(_GARBLED_RE.findall(text))
    letters = len(re.findall(r"[A-Za-z0-9]", text))
    length = max(len(text.strip()), 1)
    printable = sum(1 for ch in text if ch.isprintable() or ch in "\n\t")
    return cjk * 4.0 + letters * 0.15 + printable * 0.02 - cid * 12.0 - garbled * 15.0 - (8.0 if length < 20 else 0.0)


def _looks_garbled(text: str) -> bool:
    if not text or len(text.strip()) < 8:
        return False
    sample = text[:8000]
    n = len(sample)
    garbled = len(_GARBLED_RE.findall(sample))
    cid = len(_CID_RE.findall(sample))
    cjk = len(_CJK_RE.findall(sample))
    if garbled / n > 0.08:
        return True
    if cid >= 12 and cjk < 8:
        return True
    return False


def _normalize(text: str) -> str:
    text = text.replace("\ufeff", "").replace("\u200b", "").replace("\xa0", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CID_RE.sub(" ", text)
    text = _CONTROL_RE.sub("", text)
    lines = [" ".join(line.split()) for line in text.split("\n")]
    collapsed: list[str] = []
    blank = False
    for line in lines:
        if not line:
            if not blank:
                collapsed.append("")
            blank = True
            continue
        blank = False
        collapsed.append(line)
    return "\n".join(collapsed).strip()
